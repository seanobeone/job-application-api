from pathlib import Path
from functools import lru_cache
from docx import Document
import re

from .profile import CANDIDATE_PROFILE

STOPWORDS = {
    "the","and","or","a","an","to","of","in","for","with","on","at","as","by",
    "is","are","be","this","that","from","your","you","our","we","will","can",
    "job","role","work","team","using","use","including","experience","skills",
    "responsibilities","required","preferred","years","year","support"
}

IMPORTANT_PHRASES = [
    "cybersecurity","cyber security","information security","security operations",
    "incident response","vulnerability management","threat intelligence",
    "identity and access management","access control","network security",
    "cloud security","risk management","nist","security analyst","soc analyst",
    "aws","azure","cloud","devops","docker","linux","windows","python","fastapi",
    "api development","playwright","selenium","automation","sql",
    "technical support","desktop support","help desk","service desk",
    "systems administration","systems administrator","network support",
    "network administrator","infrastructure","endpoint","iam","grc"
]

def normalize(text):
    text = (text or "").lower()
    text = re.sub(r"[^a-z0-9+#.\- ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()

def tokens(text):
    return {
        t for t in re.findall(r"[a-z0-9+#.\-]{2,}", normalize(text))
        if t not in STOPWORDS and len(t) > 1
    }

def read_docx_text(path):
    p = Path(path)
    if not p.exists():
        return ""
    doc = Document(str(p))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                txt = cell.text.strip()
                if txt:
                    parts.append(txt)
    return "\n".join(parts)

@lru_cache(maxsize=1)
def load_resumes():
    out = {}
    for key, info in CANDIDATE_PROFILE["resume_files"].items():
        path = info["file_path"]
        text = read_docx_text(path)
        out[key] = {
            "label": info["label"],
            "file_path": path,
            "exists": Path(path).exists(),
            "text": text,
            "normalized": normalize(text),
            "tokens": tokens(text),
        }
    return out

def reload_resumes():
    load_resumes.cache_clear()
    return load_resumes()

def resume_status():
    data = load_resumes()
    return {
        k: {
            "label": v["label"],
            "file_path": v["file_path"],
            "exists": v["exists"],
            "characters_indexed": len(v["text"]),
            "unique_terms_indexed": len(v["tokens"]),
        }
        for k, v in data.items()
    }

def score_resume_against_job(resume, job_text):
    job_norm = normalize(job_text)
    job_tokens = tokens(job_text)
    resume_tokens = resume["tokens"]
    overlap = job_tokens & resume_tokens
    token_score = min(55, len(overlap) * 2)
    phrase_hits = []
    for phrase in IMPORTANT_PHRASES:
        if phrase in job_norm and phrase in resume["normalized"]:
            phrase_hits.append(phrase)
    phrase_score = min(35, len(phrase_hits) * 5)
    coverage = 0
    if job_tokens:
        coverage = min(10, round((len(overlap) / len(job_tokens)) * 40))
    score = min(100, token_score + phrase_score + coverage)
    return {"score": score, "overlap_terms": sorted(overlap)[:40], "phrase_hits": phrase_hits[:20]}

def compare_all_resumes(title, description):
    job_text = f"{title or ''}\n{description or ''}"
    resumes = load_resumes()
    scores = {}
    details = {}
    for key, resume in resumes.items():
        result = score_resume_against_job(resume, job_text)
        scores[key] = result["score"]
        details[key] = result
    best = max(scores, key=scores.get) if scores else "soc_cybersecurity"
    return {"recommended_resume": best, "resume_match_scores": scores, "resume_match_details": details}
